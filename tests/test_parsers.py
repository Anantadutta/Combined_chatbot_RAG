"""Unit tests: each of the 8 formats parses into Sections with precise locations."""
import io
import json

import pytest

from app.ingestion import ParseError, parse_bytes


def parse_file(raw_dir, name):
    return parse_bytes((raw_dir / name).read_bytes(), name)


def texts(doc):
    return [s.text for s in doc.sections]


def test_pdf_page_location_and_joined_lines(raw_dir):
    doc = parse_file(raw_dir, "HR_Policy_2026_Current.pdf")
    assert doc.file_type == "pdf"
    assert doc.sections[0].location == "page 1"
    assert "Work from home: maximum 2 days per week" in doc.full_text  # line wrap repaired


def test_docx_title_and_paragraphs(raw_dir):
    doc = parse_file(raw_dir, "Project_Orion.docx")
    assert doc.title == "Project Orion - Delivery Brief"
    locs = [s.location for s in doc.sections]
    assert locs == ["paragraph 2", "paragraph 3"]
    assert "integration due 2026-10-28" in doc.sections[1].text


def test_xlsx_every_sheet_row_is_a_record(raw_dir):
    doc = parse_file(raw_dir, "Projects_and_Budgets.xlsx")
    by_loc = {s.location: s for s in doc.sections}
    assert "sheet 'Projects', row 2" in by_loc
    change = by_loc["sheet 'Budget Changes', row 2"]
    assert change.kind == "record"
    assert change.text == "Project ID: OR-101; Change date: 2026-09-10; Approved increase USD: 10000"


def test_csv_rows_keep_column_names(raw_dir):
    doc = parse_file(raw_dir, "Employees.csv")
    assert len(doc.sections) == 6
    assert doc.sections[0].location == "row 2"
    assert "name: Maya Rao" in doc.sections[0].text


def test_json_records_have_paths(raw_dir):
    doc = parse_file(raw_dir, "Client_SLAs.json")
    cobalt = next(s for s in doc.sections if "Cobalt" in s.text)
    assert cobalt.location == "JSON path $.contracts[1]"
    assert "p1_response_minutes: 20" in cobalt.text


def test_html_sections_and_table_rows(raw_dir):
    doc = parse_file(raw_dir, "API_Guide.html")
    assert doc.title == "NexaCorp API Guide"
    row = next(s for s in doc.sections if "/tickets" in s.text)
    assert row.location == "section 'Endpoints', table 1, row 3"
    assert any("120 requests per minute" in t for t in texts(doc))


def test_markdown_heading_path_and_lines(raw_dir):
    doc = parse_file(raw_dir, "Announcements.md")
    first = doc.sections[0]
    assert first.heading == "NexaCorp announcements > 2026-09-15"
    assert first.loc["line_start"] == 4


def test_txt_line_ranges(raw_dir):
    doc = parse_file(raw_dir, "Operations_Notes.txt")
    assert doc.sections[0].location == "lines 1-5"


# ---------------------------------------------------------------- malformed input


@pytest.mark.parametrize(
    "name,data",
    [
        ("broken.pdf", b"%PDF-1.4 this is not really a pdf"),
        ("fake.pdf", b"hello world"),
        ("broken.docx", b"PK\x03\x04garbage"),
        ("fake.xlsx", b"not a zip"),
        ("bad.json", b'{"contracts": [ {"client": '),
        ("binary.txt", b"\x00\x01\x02\x03"),
        ("empty.md", b""),
        ("blank.txt", b"   \n\n  "),
        ("header_only.csv", b"a,b,c\n"),
        ("virus.exe", b"MZ..."),
    ],
)
def test_malformed_files_raise_parse_error(name, data):
    with pytest.raises(ParseError):
        parse_bytes(data, name)


def test_json_nested_and_scalar_lists():
    data = json.dumps({"name": "X", "tags": ["a", "b"], "owner": {"id": 7, "team": {"lead": "Z"}}}).encode()
    doc = parse_bytes(data, "n.json")
    locs = {s.location: s.text for s in doc.sections}
    assert locs["JSON path $"] == "name: X; tags: a, b"
    assert locs["JSON path $.owner"] == "id: 7"
    assert locs["JSON path $.owner.team"] == "lead: Z"


def test_docx_table_rows_are_records():
    import docx

    d = docx.Document()
    d.add_heading("Team", level=1)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Role"
    t.cell(1, 0).text, t.cell(1, 1).text = "Asha", "Lead"
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_bytes(buf.getvalue(), "team.docx")
    assert doc.sections[0].text == "Name: Asha; Role: Lead"
    assert doc.sections[0].location == "table 1, row 2"
    assert doc.sections[0].heading == "Team"


def test_markdown_ignores_hashes_inside_code_blocks():
    md = b"# Title\n\n```\n# not a heading\n```\n\n## Real\nBody text\n"
    doc = parse_bytes(md, "c.md")
    assert doc.sections[0].heading == "Title"
    assert "# not a heading" in doc.sections[0].text
    assert doc.sections[1].heading == "Title > Real"
