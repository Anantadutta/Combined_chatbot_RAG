"""Format-specific parsers. Each one turns raw bytes into a ParsedDocument made of Sections
that keep the document's structure (pages, headings, paragraphs, table rows, records) and a
precise location for citations.

Supported: PDF, DOCX, XLSX, CSV, JSON, HTML, TXT, Markdown.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from pathlib import Path
from typing import Any, Callable

from .models import ParsedDocument, ParseError, Section

# ---------------------------------------------------------------- helpers


def decode_text(data: bytes, filename: str) -> str:
    """Decode bytes from a text file; reject binary content."""
    if b"\x00" in data[:4096]:
        raise ParseError(f"{filename}: looks like a binary file, not text")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError(f"{filename}: could not decode text (unknown encoding)")


def clean(text: str) -> str:
    """Normalise whitespace while keeping line breaks meaningful."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def fmt_value(value: Any) -> str:
    """Format spreadsheet / JSON values for text: dates as ISO, 120000.0 as 120000."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def record_text(headers: list[str], values: list[Any]) -> str:
    """'Project ID: OR-101; Name: Orion; ...' - self-describing row text for retrieval."""
    parts = []
    for h, v in zip(headers, values):
        v = fmt_value(v)
        if v != "":
            parts.append(f"{h}: {v}")
    return "; ".join(parts)


def default_title(filename: str) -> str:
    return Path(filename).stem.replace("_", " ")


def split_paragraphs(text: str) -> list[tuple[str, int, int]]:
    """Split text into blank-line separated paragraphs -> (text, first_line, last_line), 1-based."""
    out: list[tuple[str, int, int]] = []
    buf: list[str] = []
    start = 0
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for i, line in enumerate(lines, start=1):
        if line.strip():
            if not buf:
                start = i
            buf.append(line)
        elif buf:
            out.append(("\n".join(buf), start, i - 1))
            buf = []
    if buf:
        out.append(("\n".join(buf), start, len(lines)))
    return out


def line_range(a: int, b: int) -> str:
    return f"line {a}" if a == b else f"lines {a}-{b}"


# ---------------------------------------------------------------- PDF


def parse_pdf(data: bytes, filename: str) -> ParsedDocument:
    from pypdf import PdfReader

    if not data.startswith(b"%PDF"):
        raise ParseError(f"{filename}: not a valid PDF (missing %PDF header)")
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # pypdf raises many different error types
        raise ParseError(f"{filename}: corrupt PDF ({exc})") from exc

    title = default_title(filename)
    try:
        meta_title = str(reader.metadata.title).strip() if reader.metadata and reader.metadata.title else ""
        if meta_title and meta_title.lower() not in {"(anonymous)", "untitled", "anonymous"}:
            title = meta_title
    except Exception:
        pass

    sections: list[Section] = []
    for page_no, raw in enumerate(pages, start=1):
        text = clean(raw)
        if not text:
            continue
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        for para_no, para in enumerate(paragraphs, start=1):
            # PDF text extraction wraps lines mid-sentence; join single line breaks back.
            para = re.sub(r"\s*\n\s*", " ", para)
            location = f"page {page_no}" if len(paragraphs) == 1 else f"page {page_no}, paragraph {para_no}"
            sections.append(Section(para, location, {"page": page_no, "paragraph": para_no}))
    if not sections:
        raise ParseError(f"{filename}: PDF has no extractable text (scanned image? OCR not supported)")
    return ParsedDocument(filename, "pdf", title, sections)


# ---------------------------------------------------------------- DOCX


def parse_docx(data: bytes, filename: str) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    if not data.startswith(b"PK"):
        raise ParseError(f"{filename}: not a valid DOCX (not a zip container)")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"{filename}: corrupt DOCX ({exc})") from exc

    title = ""
    heading: str | None = None
    sections: list[Section] = []
    para_no = 0
    table_no = 0
    # Walk the body in document order so paragraphs and tables keep their sequence.
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(element, document)
            text = clean(para.text)
            if not text:
                continue
            para_no += 1
            style = (para.style.name if para.style is not None else "") or ""
            if style == "Title":
                title = title or text
                heading = text
                continue
            if style.startswith("Heading"):
                heading = text
                continue
            sections.append(
                Section(text, f"paragraph {para_no}", {"paragraph": para_no}, heading=heading)
            )
        elif tag == "tbl":
            table_no += 1
            rows = [[clean(c.text) for c in r.cells] for r in Table(element, document).rows]
            if not rows:
                continue
            headers = rows[0]
            for row_no, row in enumerate(rows[1:], start=2):
                text = record_text(headers, row)
                if text:
                    sections.append(
                        Section(
                            text,
                            f"table {table_no}, row {row_no}",
                            {"table": table_no, "row": row_no},
                            heading=heading,
                            kind="record",
                        )
                    )
    if not sections:
        raise ParseError(f"{filename}: DOCX contains no text")
    return ParsedDocument(filename, "docx", title or default_title(filename), sections)


# ---------------------------------------------------------------- XLSX


def parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    import openpyxl

    if not data.startswith(b"PK"):
        raise ParseError(f"{filename}: not a valid XLSX (not a zip container)")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"{filename}: corrupt XLSX ({exc})") from exc

    sections: list[Section] = []
    for sheet in workbook.worksheets:
        headers: list[str] | None = None
        for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = list(row)
            if all(fmt_value(v) == "" for v in values):
                continue
            if headers is None:  # first non-empty row = header row
                headers = [fmt_value(v) or f"Column {i + 1}" for i, v in enumerate(values)]
                continue
            text = record_text(headers, values)
            sections.append(
                Section(
                    text,
                    f"sheet '{sheet.title}', row {row_no}",
                    {"sheet": sheet.title, "row": row_no},
                    heading=f"Sheet: {sheet.title}",
                    kind="record",
                )
            )
    workbook.close()
    if not sections:
        raise ParseError(f"{filename}: workbook has no data rows")
    return ParsedDocument(filename, "xlsx", default_title(filename), sections)


# ---------------------------------------------------------------- CSV


def parse_csv(data: bytes, filename: str) -> ParsedDocument:
    text = decode_text(data, filename)
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error as exc:
        raise ParseError(f"{filename}: malformed CSV ({exc})") from exc
    rows = [r for r in rows if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise ParseError(f"{filename}: CSV needs a header row and at least one data row")

    headers = [h.strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    sections = []
    for row_no, row in enumerate(rows[1:], start=2):
        sections.append(
            Section(record_text(headers, row), f"row {row_no}", {"row": row_no}, kind="record")
        )
    return ParsedDocument(filename, "csv", default_title(filename), sections)


# ---------------------------------------------------------------- JSON


def _is_scalar(v: Any) -> bool:
    return v is None or isinstance(v, (str, int, float, bool))


def parse_json(data: bytes, filename: str) -> ParsedDocument:
    try:
        obj = json.loads(decode_text(data, filename))
    except json.JSONDecodeError as exc:
        raise ParseError(f"{filename}: invalid JSON ({exc})") from exc

    sections: list[Section] = []

    def emit(path: str, pairs: list[tuple[str, Any]]) -> None:
        text = "; ".join(f"{k}: {fmt_value(v)}" for k, v in pairs if fmt_value(v) != "")
        if text:
            sections.append(
                Section(text, f"JSON path {path}", {"json_path": path}, heading=path, kind="record")
            )

    def walk(node: Any, path: str) -> None:
        if isinstance(node, dict):
            scalars = [(k, v) for k, v in node.items() if _is_scalar(v)]
            scalars += [
                (k, ", ".join(fmt_value(x) for x in v))
                for k, v in node.items()
                if isinstance(v, list) and v and all(_is_scalar(x) for x in v)
            ]
            emit(path, scalars)  # the flat fields of this object become one record
            for k, v in node.items():
                if isinstance(v, dict) or (isinstance(v, list) and not all(_is_scalar(x) for x in v)):
                    walk(v, f"{path}.{k}")
        elif isinstance(node, list):
            for i, item in enumerate(node):
                if _is_scalar(item):
                    emit(f"{path}[{i}]", [("value", item)])
                else:
                    walk(item, f"{path}[{i}]")
        else:
            emit(path, [("value", node)])

    walk(obj, "$")
    if not sections:
        raise ParseError(f"{filename}: JSON contains no data")
    return ParsedDocument(filename, "json", default_title(filename), sections)


# ---------------------------------------------------------------- HTML


def parse_html(data: bytes, filename: str) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(decode_text(data, filename), "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()

    title = clean(soup.title.get_text()) if soup.title and soup.title.get_text().strip() else ""
    heading: str | None = None
    sections: list[Section] = []
    counters = {"paragraph": 0, "table": 0}
    block_tags = ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote", "table"]

    for el in soup.find_all(block_tags):
        # skip blocks nested inside another handled block (e.g. <p> inside <li> or <table>)
        if el.find_parent(["table", "li", "blockquote"]) is not None:
            continue
        name = el.name
        if name.startswith("h") and len(name) == 2:
            heading = clean(el.get_text(" "))
            title = title or heading
            continue
        if name == "table":
            counters["table"] += 1
            rows = el.find_all("tr")
            headers = [clean(c.get_text(" ")) for c in rows[0].find_all(["th", "td"])] if rows else []
            for row_no, tr in enumerate(rows[1:], start=2):
                cells = [clean(c.get_text(" ")) for c in tr.find_all(["th", "td"])]
                text = record_text(headers, cells)
                if text:
                    loc = f"table {counters['table']}, row {row_no}"
                    sections.append(
                        Section(
                            text,
                            f"section '{heading}', {loc}" if heading else loc,
                            {"section": heading, "table": counters["table"], "row": row_no},
                            heading=heading,
                            kind="record",
                        )
                    )
            continue
        text = clean(el.get_text(" "))
        if not text:
            continue
        counters["paragraph"] += 1
        loc = f"paragraph {counters['paragraph']}"
        sections.append(
            Section(
                text,
                f"section '{heading}', {loc}" if heading else loc,
                {"section": heading, "paragraph": counters["paragraph"]},
                heading=heading,
            )
        )

    if not sections:  # page without block tags: fall back to plain text
        body_text = clean(soup.get_text("\n"))
        if body_text:
            sections.append(Section(body_text, "body", {}))
    if not sections:
        raise ParseError(f"{filename}: HTML contains no text")
    return ParsedDocument(filename, "html", title or default_title(filename), sections)


# ---------------------------------------------------------------- Markdown


def parse_markdown(data: bytes, filename: str) -> ParsedDocument:
    text = decode_text(data, filename)
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")

    title = ""
    heading_stack: list[tuple[int, str]] = []
    sections: list[Section] = []
    buf: list[str] = []
    buf_start = 0
    in_code = False

    def heading_path() -> str | None:
        return " > ".join(h for _, h in heading_stack) or None

    def flush(end_line: int) -> None:
        nonlocal buf
        body = clean("\n".join(buf))
        if body:
            path = heading_path()
            loc = line_range(buf_start, end_line)
            sections.append(
                Section(
                    body,
                    f"section '{path}', {loc}" if path else loc,
                    {"section": path, "line_start": buf_start, "line_end": end_line},
                    heading=path,
                )
            )
        buf = []

    for i, line in enumerate(lines, start=1):
        if line.strip().startswith("```"):
            in_code = not in_code
        match = None if in_code else re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", line)
        if match:
            flush(i - 1)
            level, name = len(match.group(1)), match.group(2)
            if level == 1 and not title:
                title = name
            heading_stack = [(lv, h) for lv, h in heading_stack if lv < level] + [(level, name)]
            continue
        if not line.strip() and not in_code:
            flush(i - 1)
            continue
        if not buf:
            buf_start = i
        buf.append(line)
    flush(len(lines))

    if not sections:
        raise ParseError(f"{filename}: Markdown file has no body text")
    return ParsedDocument(filename, "md", title or default_title(filename), sections)


# ---------------------------------------------------------------- TXT


def parse_txt(data: bytes, filename: str) -> ParsedDocument:
    text = decode_text(data, filename)
    sections = [
        Section(clean(para), line_range(a, b), {"line_start": a, "line_end": b})
        for para, a, b in split_paragraphs(text)
    ]
    if not sections:
        raise ParseError(f"{filename}: text file is empty")
    return ParsedDocument(filename, "txt", default_title(filename), sections)


# ---------------------------------------------------------------- registry

PARSERS: dict[str, Callable[[bytes, str], ParsedDocument]] = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".csv": parse_csv,
    ".json": parse_json,
    ".html": parse_html,
    ".htm": parse_html,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".txt": parse_txt,
}

SUPPORTED_EXTENSIONS = frozenset(PARSERS)


def parse_bytes(data: bytes, filename: str) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ParseError(f"{filename}: unsupported file type '{ext or '(none)'}'")
    if not data:
        raise ParseError(f"{filename}: file is empty")
    return parser(data, filename)
